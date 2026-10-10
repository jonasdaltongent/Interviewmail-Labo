#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert het werkdocument voor les 3 "De formele e-mail" (Onderzoekslabo Jongeren & Welzijn, klas 3MWWE):

  Interviewmail.docx   het werkdocument dat de leerling INLEVERT (wordt in Classroom een Google-document)

Gebruik:  python3 maak_werkdocumenten.py     (vereist python-docx)

Waarom geen zip en geen Gmail: de leerlingen hebben met hun schoolaccount geen Gmail (Jonas, 09-10-2026 en
10-10-2026: "Werkdocument zoals 3MWb"), en per les gebruiken ze één werkplek naast Classroom en de lespagina
(_afspraken/didactiek.md). De e-mail komt daarom in dit document, in een venster dat eruitziet als "Nieuw
bericht" in Gmail. Model: Stagemail van MW les 3.

LET OP bij aanpassen:
 - Deel 1, rij 1 (het onderwerp) is het uitgewerkte voorbeeld: dat doet de leraar voor (dia 6). Laat het ingevuld.
 - De adressen eindigen op .example: dat domein bestaat nooit echt (RFC 2606). Geen echte organisatie.
 - Alle namen en gegevens zijn fictief. Noors e-mail gaat bewust over een ander onderzoek en aan iemand anders
   dan de e-mail van de leerling: de leerling neemt de opbouw over, niet de zinnen.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x2A, 0x39, 0x73)
GREY = RGBColor(0x5F, 0x63, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HERE = os.path.dirname(os.path.abspath(__file__))

JONGERENWERKER = ("mevrouw Hanne De Clercq", "hanne.declercq@delantaarn.example")
BEGELEIDER = ("mevrouw Ines Coppens", "ines.coppens@onderzoekslabo.example")


# ---------- hulpfuncties ----------
def basis_document(marge=2.0):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(marge)
        s.left_margin = s.right_margin = Cm(marge)
    return doc


def kop(doc, tekst, grootte=16, ruimte_voor=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(ruimte_voor)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(tekst)
    r.bold = True
    r.font.size = Pt(grootte)
    r.font.color.rgb = NAVY
    return p


def tekst(doc, s, cursief=False, klein=False, vet=False, na=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(na)
    r = p.add_run(s)
    r.italic = cursief
    r.bold = vet
    if klein:
        r.font.size = Pt(9.5)
        r.font.color.rgb = GREY
    return p


def schaduw(cel, kleur="E4E8F6"):
    tcPr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), kleur)
    tcPr.append(shd)


def randen(tabel, buiten="9AA0A6", binnen=None):
    """Randen van een tabel: alleen een kader rondom, of ook lijnen binnenin."""
    tblPr = tabel._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for kant in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + kant)
        kleur = buiten if kant in ("top", "left", "bottom", "right") else binnen
        if kleur:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "8")
            e.set(qn("w:color"), kleur)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def onderrand(cel, kleur="DADCE0"):
    tcPr = cel._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), "6")
    e.set(qn("w:color"), kleur)
    b.append(e)
    tcPr.append(b)


def breedtes(tabel, cm):
    """Vaste kolombreedtes, zodat Google Documenten de tabel niet zelf herschikt."""
    tabel.autofit = False
    for kol, w in zip(tabel.columns, cm):
        kol.width = Cm(w)
    for rij in tabel.rows:
        for cel, w in zip(rij.cells, cm):
            cel.width = Cm(w)


def celtekst(cel, s, vet=False, kleur=None, grootte=None):
    cel.text = ""
    r = cel.paragraphs[0].add_run(s)
    r.bold = vet
    if kleur:
        r.font.color.rgb = kleur
    if grootte:
        r.font.size = Pt(grootte)
    return r


def hoogte(rij, cm):
    trPr = rij._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def antwoordlijnen(doc, aantal=2):
    for _ in range(aantal):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run("_" * 70)


def venster(doc, delen):
    """Een venster dat eruitziet als 'Nieuw bericht' in Gmail.
    delen: (grijs label links van een tekstvak in het bericht, hoogte van het vak in cm)."""
    rijen = 1 + 4 + len(delen) + 1
    t = doc.add_table(rows=rijen, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    randen(t, buiten="9AA0A6")
    breedtes(t, [3.2, 13.8])

    kopcel = t.rows[0].cells[0].merge(t.rows[0].cells[1])
    schaduw(kopcel, "F2F6FC")
    celtekst(kopcel, "Nieuw bericht", vet=True, kleur=RGBColor(0x20, 0x21, 0x24))

    for i, veld in enumerate(["Aan", "Cc", "Bcc", "Onderwerp"], start=1):
        links, rechts = t.rows[i].cells
        celtekst(links, veld, kleur=GREY)
        onderrand(links)
        onderrand(rechts)
        hoogte(t.rows[i], 0.8)

    for j, (label, h) in enumerate(delen):
        rij = t.rows[5 + j]
        celtekst(rij.cells[0], label, kleur=GREY, grootte=8.5)
        hoogte(rij, h)

    voet = t.rows[-1]
    knop = voet.cells[0]
    schaduw(knop, "0B57D0")
    celtekst(knop, "Verzenden", vet=True, kleur=WHITE)
    knop.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    celtekst(voet.cells[1], "Niet echt: je levert dit document in, in Classroom.", kleur=GREY, grootte=9)
    return t



# ---------- de twee voorbeeldmails ----------
SENNE = {
    "naam": "Senne",
    "aan": "hanne.declercq@delantaarn.example; ines.coppens@onderzoekslabo.example",
    "cc": "",
    "onderwerp": "vraagje",
    "tekst": [
        "hallo,",
        "ik zit in het 3de jaar en we moeten een interview doen voor een taak over slaap. "
        "kunt ge morgen na school?",
        "thx!! 🙂",
        "Senne",
    ],
}
NOOR = {
    "naam": "Noor",
    "aan": "bram.willems@clb-ketel.example",
    "cc": "ines.coppens@onderzoekslabo.example",
    "onderwerp": "Vraag om een interview over groepschats – Onderzoekslabo",
    "tekst": [
        "Beste meneer Willems,",
        "Mijn naam is Noor Van Hecke en ik zit in 3MWWE van GO! Dalton Gent. Samen met mijn onderzoeksgroep "
        "onderzoek ik namens het Onderzoekslabo Jongeren & Welzijn groepsdruk in groepschats.",
        "Mag ik u daarover interviewen? Het interview duurt ongeveer 20 minuten. Het kan op woensdag 28 oktober "
        "om 15 uur of op donderdag 29 oktober om 13 uur, bij het CLB of online. Uw naam komt niet in ons "
        "verslag. We gebruiken uw antwoorden alleen voor dit schoolonderzoek.",
        "Kan u mij vóór dinsdag 20 oktober laten weten welk moment past? Alvast bedankt voor uw tijd.",
        "Met vriendelijke groeten,",
        "Noor Van Hecke\n3MWWE\nOnderzoekslabo Jongeren & Welzijn\nGO! Dalton Gent",
    ],
}


def voorbeeldmail(cel, mail):
    """Zet een korte nagemaakte e-mail in een cel van een tabel."""
    cel.text = ""
    kopje = cel.paragraphs[0]
    r = kopje.add_run("E-mail van " + mail["naam"])
    r.bold = True
    r.font.color.rgb = NAVY
    r.font.size = Pt(10.5)
    kopje.paragraph_format.space_after = Pt(3)
    for label, waarde in (("Aan", mail["aan"]), ("Cc", mail["cc"] or "—"), ("Onderwerp", mail["onderwerp"])):
        p = cel.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        a = p.add_run(label + ":  ")
        a.font.size = Pt(9)
        a.font.color.rgb = GREY
        b = p.add_run(waarde)
        b.font.size = Pt(9)
    for s in mail["tekst"]:
        p = cel.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(s)
        rr.font.size = Pt(9.5)


# ---------- het werkdocument ----------
def interviewmail():
    doc = basis_document()

    p = doc.add_paragraph()
    r = p.add_run("INTERVIEWMAIL")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY
    tekst(doc, "Les 3 — De formele e-mail  ·  Toegepaste Informatica  ·  Onderzoekslabo Jongeren & Welzijn",
          klein=True, na=10)

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, label in enumerate(["Naam:", "Klas:", "Datum:"]):
        c = t.rows[0].cells[i]
        c.text = label + " "
        c.paragraphs[0].runs[0].bold = True
    breedtes(t, [7.0, 4.0, 6.0])
    doc.add_paragraph()

    kop(doc, "Zo werk je", 13, ruimte_voor=6)
    for s in [
        "1.  Op de lespagina lees je wat je moet doen, stap voor stap.",
        "2.  Je schrijft alles in DIT document. Je hebt niets anders nodig.",
        "3.  Je e-mail schrijf je in deel 3, in een venster dat eruitziet als Nieuw bericht in Gmail.",
        "4.  Klaar? Klik in de opdracht op Inleveren. Er wordt niets echt verstuurd.",
    ]:
        p = doc.add_paragraph(s)
        p.paragraph_format.space_after = Pt(2)

    # ---- Je onderzoek (het kader met de feiten)
    kop(doc, "Je onderzoek", 13)
    k = doc.add_table(rows=1, cols=1)
    randen(k, buiten="A7B1DE")
    breedtes(k, [17.0])
    cel = k.rows[0].cells[0]
    schaduw(cel, "EEF0F9")
    cel.text = ""
    regels = [
        ("Je onderzoeksgroep in het Onderzoekslabo Jongeren & Welzijn onderzoekt schermgebruik en slaap bij "
         "jongeren (het onderzoeksplan uit les 2). Jullie willen een interview met een jongerenwerker.", False),
        ("Het interview:  " + JONGERENWERKER[0] + ", jongerenwerker bij Jongerenhuis De Lantaarn  —  "
         + JONGERENWERKER[1], True),
        ("Je onderzoeksbegeleider:  " + BEGELEIDER[0] + "  —  " + BEGELEIDER[1], True),
        ("Wat je vraagt:  een interview van ongeveer 20 minuten  ·  twee mogelijke momenten: woensdag "
         "21 oktober om 14 uur of donderdag 22 oktober om 16 uur  ·  in het jongerenhuis of online.", False),
        ("Wat er met haar antwoorden gebeurt:  ze worden gebruikt voor een schoolonderzoek en haar naam komt "
         "niet in het verslag.", False),
        ("Tegen wanneer je een antwoord hoopt:  maandag 19 oktober.", False),
        ("Je schrijft zelf de zinnen. Alle namen en adressen zijn verzonnen.", None),
    ]
    for n, (s, vet) in enumerate(regels):
        pp = cel.paragraphs[0] if n == 0 else cel.add_paragraph()
        pp.paragraph_format.space_after = Pt(4)
        rr = pp.add_run(s)
        if vet:
            rr.bold = True
        if vet is None:
            rr.italic = True
            rr.font.size = Pt(9)
            rr.font.color.rgb = GREY

    # ---- Deel 1
    kop(doc, "Deel 1 — Goed of zwak?")
    tekst(doc, "Senne en Noor vroegen allebei een interview aan. Lees hun e-mails. Noor schreef over een ander "
               "onderzoek en aan iemand anders: kijk naar hoe ze het opbouwt, niet naar haar zinnen.", klein=True)
    mt = doc.add_table(rows=1, cols=2)
    mt.style = "Table Grid"
    mt.alignment = WD_TABLE_ALIGNMENT.CENTER
    voorbeeldmail(mt.rows[0].cells[0], SENNE)
    voorbeeldmail(mt.rows[0].cells[1], NOOR)
    breedtes(mt, [6.2, 10.8])
    doc.add_paragraph()
    tekst(doc, "Vergelijk ze per onderdeel. Rij 1 deed je leraar voor. Schrijf bij rij 2 tot 6 in een paar "
               "woorden wat Senne deed en wat Noor deed. Leg dan uit waarom het bij Noor beter werkt.", klein=True)
    t = doc.add_table(rows=7, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Onderdeel", "Senne", "Noor", "Waarom werkt het bij Noor beter?"]):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        schaduw(c)
    rijen = [
        ("1  onderwerp", "vraagje", "Vraag om een interview over groepschats – Onderzoekslabo",
         "Bij vraagje weet de ontvanger niet waarover het gaat. Bij Noor zie je waarover (een interview over "
         "groepschats) en van wie (Onderzoekslabo) nog voor je de e-mail opent."),
        ("2  aanspreking", "", "", ""),
        ("3  inleiding", "", "", ""),
        ("4  midden: de vraag", "", "", ""),
        ("5  slot", "", "", ""),
        ("6  slotgroet en handtekening", "", "", ""),
    ]
    for i, (onderdeel, senne, noor, waarom) in enumerate(rijen, start=1):
        cells = t.rows[i].cells
        cells[0].text = onderdeel
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = senne
        cells[2].text = noor
        cells[3].text = waarom
        hoogte(t.rows[i], 1.0 if i == 1 else 1.7)
    for cel_ in t.rows[1].cells[1:]:
        for run in cel_.paragraphs[0].runs:
            run.font.size = Pt(9.5)
    breedtes(t, [3.0, 3.4, 4.4, 6.2])
    doc.add_paragraph()

    # ---- Deel 2
    kop(doc, "Deel 2 — Aan, Cc of Bcc?")
    tekst(doc, "Schrijf bij elke situatie Aan, Cc of Bcc.", klein=True)
    t = doc.add_table(rows=5, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Situatie", "Aan, Cc of Bcc?"]):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        schaduw(c)
    situaties = [
        "a.  Je mailt de jongerenwerker die je wil interviewen.",
        "b.  Je onderzoeksbegeleider wil je e-mail lezen, maar moet zelf niets doen.",
        "c.  Je nodigt 24 leerlingen van het 4de jaar uit voor je enquête. Ze mogen elkaars e-mailadres niet zien.",
        "d.  Je twee groepsgenoten moeten weten wat je vroeg.",
    ]
    for i, s in enumerate(situaties, start=1):
        t.rows[i].cells[0].text = s
        hoogte(t.rows[i], 0.9)
    breedtes(t, [12.5, 4.5])
    doc.add_paragraph()
    tekst(doc, "Vraag 1. Waarom zet je de 24 leerlingen niet in Cc?", vet=True)
    antwoordlijnen(doc, 2)

    # ---- Deel 3
    kop(doc, "Deel 3 — Mijn interviewaanvraag")
    tekst(doc, "Schrijf de e-mail die Senne had moeten sturen, maar met jouw naam. Klik in een vak en typ. De "
               "grijze woorden links zeggen wat in elk vak hoort. Laat Bcc leeg.", klein=True)
    venster(doc, [("aanspreking", 0.9),
                  ("inleiding:\nwie ben je, namens wie, waarom mail je?", 2.0),
                  ("midden:\nwat vraag je?", 3.0),
                  ("slot:\ntegen wanneer, bedanken", 1.6),
                  ("slotgroet", 0.9),
                  ("handtekening:\nnaam, klas, labo, school", 2.4)])
    doc.add_paragraph()

    # ---- Deel 4
    kop(doc, "Deel 4 — Feedback van je buur")
    tekst(doc, "Je buur leest je e-mail op jouw scherm en kruist aan wat klopt. Klopt iets nog niet? Pas je "
               "e-mail aan.", klein=True)
    for s in [
        "Bij Aan staat de jongerenwerker, bij Cc de begeleider. Bcc is leeg.",
        "Het onderwerp zegt waarover de e-mail gaat en van wie hij komt.",
        "De aanspreking is Beste mevrouw De Clercq, (of Geachte), met een komma.",
        "Overal staat u en uw, nergens je of jij.",
        "In de inleiding staat wie ze is, namens wie ze mailt en waarom.",
        "In het midden staat wat ze vraagt: hoe lang, wanneer, waar, en wat er met de antwoorden gebeurt.",
        "In het slot staat tegen wanneer ze een antwoord hoopt, en ze bedankt.",
        "De slotgroet is Met vriendelijke groeten, met daaronder naam, klas, Onderzoekslabo en school. "
        "Geen emoji's of chattaal.",
    ]:
        p = doc.add_paragraph("☐  " + s)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph()
    tekst(doc, "Dit is goed (een top van je buur):", vet=True)
    antwoordlijnen(doc, 1)
    tekst(doc, "Dit kan beter (een tip van je buur):", vet=True)
    antwoordlijnen(doc, 1)
    tekst(doc, "Wat heb ik aangepast?", vet=True)
    antwoordlijnen(doc, 1)

    # ---- Deel 5
    kop(doc, "Deel 5 — Slotvraag")
    tekst(doc, "Vraag 2. Welke twee dingen moet mevrouw De Clercq na het lezen van je e-mail zeker weten?",
          vet=True)
    antwoordlijnen(doc, 2)
    doc.add_paragraph()
    tekst(doc, "Klaar? Klik in de opdracht op Inleveren.", vet=True)

    doc.save(os.path.join(HERE, "Interviewmail.docx"))


if __name__ == "__main__":
    interviewmail()
    print("Gemaakt: Interviewmail.docx")
